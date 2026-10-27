"""Tests para el handler de Word."""
from docx import Document
from mi_automatizador.core import word_handler


def test_convertir_genera_docx(excel_temporal, carpeta_output_temporal, monkeypatch):
    """Verifica que se genera un .docx válido."""
    from mi_automatizador.core import common
    monkeypatch.setattr(common, "CARPETA_OUTPUT", carpeta_output_temporal)

    resultado = word_handler.convertir(excel_temporal)

    assert resultado is not None
    assert resultado.exists()
    assert resultado.suffix == ".docx"


def test_contiene_tabla(excel_temporal, carpeta_output_temporal, monkeypatch):
    """Verifica que el Word contiene una tabla con los datos."""
    from mi_automatizador.core import common
    monkeypatch.setattr(common, "CARPETA_OUTPUT", carpeta_output_temporal)

    resultado = word_handler.convertir(excel_temporal)

    doc = Document(resultado)
    assert len(doc.tables) == 1
    tabla = doc.tables[0]
    assert len(tabla.rows) == 4  # encabezado + 3 filas
    assert tabla.cell(0, 0).text == "Nombre"
    assert tabla.cell(1, 0).text == "Ana"


def test_excel_vacio_retorna_none(excel_vacio, carpeta_output_temporal, monkeypatch):
    """Un Excel sin filas debe retornar None (no genera Word)."""
    from mi_automatizador.core import common
    monkeypatch.setattr(common, "CARPETA_OUTPUT", carpeta_output_temporal)

    resultado = word_handler.convertir(excel_vacio)
    assert resultado is None